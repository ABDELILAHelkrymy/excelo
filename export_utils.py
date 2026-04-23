import pathlib


def ensure_excel_output_path(path: str) -> str:
    """Normalize export paths so every save targets an .xlsx file."""
    normalized = path.strip()
    if not normalized:
        return ""

    output_path = pathlib.Path(normalized)
    if output_path.suffix.lower() == ".xlsx":
        return str(output_path)
    if output_path.suffix:
        output_path = output_path.with_suffix(".xlsx")
    else:
        output_path = output_path.with_name(f"{output_path.name}.xlsx")
    return str(output_path)

def apply_zebra(path: str):
    """Apply zebra striping (alternating row colors + styled header) to an xlsx file."""
    import os
    import logging
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import PatternFill, Font
        wb = load_workbook(path)
        fill_hdr  = PatternFill(fill_type="solid", fgColor="4472C4")
        fill_even = PatternFill(fill_type="solid", fgColor="DCE6F1")
        fill_odd  = PatternFill(fill_type="solid", fgColor="FFFFFF")
        font_hdr  = Font(bold=True, color="FFFFFF")
        for ws in wb.worksheets:
            ncols = ws.max_column or 1
            nrows = ws.max_row or 1
            for c in range(1, ncols + 1):
                cell = ws.cell(1, c)
                cell.fill = fill_hdr
                cell.font = font_hdr
            for r in range(2, nrows + 1):
                rf = fill_even if (r % 2 == 0) else fill_odd
                for c in range(1, ncols + 1):
                    ws.cell(r, c).fill = rf
        tmp_path = path + ".tmp"
        wb.save(tmp_path)
        os.replace(tmp_path, path)
    except Exception as e:
        logging.exception(f"Erreur lors de l'application du style zébré sur {path}")
