from __future__ import annotations

from copy import copy as copy_style
import math
from pathlib import Path
from typing import Callable, Iterable

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .file_service import load_tabular_file


ProgressCallback = Callable[[int, int, str], None]


def split_workbook_sheets(
    file_paths: Iterable[str],
    *,
    pattern: str,
    output_format: str,
    output_dir: str | None,
) -> int:
    total_written = 0
    for file_path in file_paths:
        source = Path(file_path)
        with pd.ExcelFile(source) as excel_file:
            if len(excel_file.sheet_names) < 2:
                continue
            destination = Path(output_dir) if output_dir else source.parent
            destination.mkdir(parents=True, exist_ok=True)
            for sheet_name in excel_file.sheet_names:
                df = excel_file.parse(sheet_name)
                safe_name = _sanitize_token(
                    pattern.format(fichier=source.stem, onglet=sheet_name),
                    max_length=180,
                )
                if output_format == "csv":
                    target = destination / f"{safe_name}.csv"
                    df.to_csv(target, index=False, encoding="utf-8-sig")
                else:
                    target = destination / f"{safe_name}.xlsx"
                    df.to_excel(target, index=False, engine="xlsxwriter")
                total_written += 1
    return total_written


def split_by_columns(
    *,
    source_path: str,
    dataframe: pd.DataFrame,
    split_columns: list[str],
    keep_columns: list[str],
    output_dir: str | None,
    zebra: bool,
    all_sheets: bool,
    sheet_names: list[str] | None,
    selected_sheet: str | None,
    progress_callback: ProgressCallback | None = None,
) -> int:
    source = Path(source_path)
    destination = Path(output_dir) if output_dir else source.parent
    destination.mkdir(parents=True, exist_ok=True)
    is_csv = source.suffix.lower() == ".csv"

    if not split_columns:
        raise ValueError("Ajoutez au moins une colonne de découpe.")
    if not keep_columns:
        raise ValueError("Cochez au moins une colonne à conserver.")

    if is_csv:
        total_groups = _count_groups(dataframe, split_columns)
    elif all_sheets and sheet_names:
        with pd.ExcelFile(source) as excel_file:
            total_groups = sum(_count_groups(excel_file.parse(sheet_name), split_columns) for sheet_name in sheet_names)
    else:
        total_groups = _count_groups(dataframe, split_columns)

    if total_groups == 0:
        raise ValueError("Aucun fichier à créer avec cette configuration.")

    _notify(progress_callback, 0, total_groups, "Préparation …")

    if is_csv:
        return _split_csv_groups(
            dataframe=dataframe,
            split_columns=split_columns,
            keep_columns=keep_columns,
            destination=destination,
            total_groups=total_groups,
            progress_callback=progress_callback,
        )

    workbook = load_workbook(source)
    try:
        with pd.ExcelFile(source) as excel_file:
            written = 0
            if all_sheets and sheet_names:
                for sheet_name in sheet_names:
                    if sheet_name not in workbook.sheetnames:
                        continue
                    worksheet = workbook[sheet_name]
                    sheet_df = excel_file.parse(sheet_name)
                    sheet_destination = destination / _sanitize_token(sheet_name, max_length=60)
                    written = _split_worksheet_groups(
                        worksheet=worksheet,
                        dataframe=sheet_df,
                        split_columns=split_columns,
                        keep_columns=keep_columns,
                        destination=sheet_destination,
                        zebra=zebra,
                        written=written,
                        total_groups=total_groups,
                        progress_callback=progress_callback,
                    )
                return written

            worksheet = workbook[selected_sheet] if selected_sheet and selected_sheet in workbook.sheetnames else workbook.active
            return _split_worksheet_groups(
                worksheet=worksheet,
                dataframe=dataframe,
                split_columns=split_columns,
                keep_columns=keep_columns,
                destination=destination,
                zebra=zebra,
                written=0,
                total_groups=total_groups,
                progress_callback=progress_callback,
            )
    finally:
        workbook.close()


def _split_csv_groups(
    *,
    dataframe: pd.DataFrame,
    split_columns: list[str],
    keep_columns: list[str],
    destination: Path,
    total_groups: int,
    progress_callback: ProgressCallback | None,
) -> int:
    destination.mkdir(parents=True, exist_ok=True)
    keep_present = [column for column in keep_columns if column in dataframe.columns]
    split_present = [column for column in split_columns if column in dataframe.columns]
    written = 0
    for keys, group_df in dataframe.groupby(split_present, dropna=False):
        target = destination / f"{_safe_group_name(keys)}.xlsx"
        group_df.loc[:, keep_present].to_excel(target, index=False, engine="xlsxwriter")
        written += 1
        _notify(progress_callback, written, total_groups, f"{written} / {total_groups} fichier(s) créé(s)")
    return written


def _split_worksheet_groups(
    *,
    worksheet,
    dataframe: pd.DataFrame,
    split_columns: list[str],
    keep_columns: list[str],
    destination: Path,
    zebra: bool,
    written: int,
    total_groups: int,
    progress_callback: ProgressCallback | None,
) -> int:
    destination.mkdir(parents=True, exist_ok=True)
    header_map = {}
    for column_index in range(1, (worksheet.max_column or 0) + 1):
        value = worksheet.cell(1, column_index).value
        if value is not None:
            header_map[str(value)] = column_index

    keep_indices = [header_map[column] for column in keep_columns if column in header_map]
    split_present = [column for column in split_columns if column in dataframe.columns]
    if not keep_indices or not split_present:
        return written

    row_cache = _build_row_cache(worksheet, keep_indices)
    width_cache = _build_width_cache(worksheet, keep_indices)
    header_height = worksheet.row_dimensions[1].height if 1 in worksheet.row_dimensions else None

    header_fill = PatternFill(fill_type="solid", fgColor="1F3864") if zebra else None
    even_fill = PatternFill(fill_type="solid", fgColor="DCE6F1") if zebra else None
    odd_fill = PatternFill(fill_type="solid", fgColor="FFFFFF") if zebra else None
    header_font = Font(bold=True, color="FFFFFF") if zebra else None

    for keys, group_df in dataframe.groupby(split_present, dropna=False):
        target = destination / f"{_safe_group_name(keys)}.xlsx"
        workbook = Workbook()
        out_sheet = workbook.active

        for new_column_index, source_column_index in enumerate(keep_indices, start=1):
            _copy_cell(row_cache[1][source_column_index], out_sheet.cell(1, new_column_index))
            if source_column_index in width_cache:
                out_sheet.column_dimensions[get_column_letter(new_column_index)].width = width_cache[source_column_index]
            if zebra and header_fill and header_font:
                header_cell = out_sheet.cell(1, new_column_index)
                header_cell.fill = header_fill
                header_cell.font = header_font
        if header_height is not None:
            out_sheet.row_dimensions[1].height = header_height

        target_row = 2
        for frame_index in group_df.index:
            source_row = int(frame_index) + 2
            source_cells = row_cache.get(source_row)
            if source_cells is None:
                source_cells = {column_index: worksheet.cell(source_row, column_index) for column_index in keep_indices}
            for new_column_index, source_column_index in enumerate(keep_indices, start=1):
                _copy_cell(source_cells[source_column_index], out_sheet.cell(target_row, new_column_index))
                if zebra and even_fill and odd_fill:
                    out_sheet.cell(target_row, new_column_index).fill = even_fill if target_row % 2 == 0 else odd_fill
            if source_row in worksheet.row_dimensions:
                out_sheet.row_dimensions[target_row].height = worksheet.row_dimensions[source_row].height
            target_row += 1

        workbook.save(target)
        workbook.close()
        written += 1
        _notify(progress_callback, written, total_groups, f"{written} / {total_groups} fichier(s) créé(s)")

    return written


def _count_groups(dataframe: pd.DataFrame, split_columns: list[str]) -> int:
    split_present = [column for column in split_columns if column in dataframe.columns]
    if not split_present:
        return 0
    return dataframe.groupby(split_present, dropna=False).ngroups


def _build_row_cache(worksheet, keep_indices: list[int]) -> dict[int, dict[int, object]]:
    row_cache: dict[int, dict[int, object]] = {}
    for row_index in range(1, (worksheet.max_row or 0) + 1):
        row_cache[row_index] = {column_index: worksheet.cell(row_index, column_index) for column_index in keep_indices}
    return row_cache


def _build_width_cache(worksheet, keep_indices: list[int]) -> dict[int, float | None]:
    widths: dict[int, float | None] = {}
    for source_column_index in keep_indices:
        letter = get_column_letter(source_column_index)
        if letter in worksheet.column_dimensions:
            widths[source_column_index] = worksheet.column_dimensions[letter].width
    return widths


def _copy_cell(source_cell, target_cell) -> None:
    target_cell.value = source_cell.value
    if source_cell.has_style:
        target_cell.font = copy_style(source_cell.font)
        target_cell.fill = copy_style(source_cell.fill)
        target_cell.border = copy_style(source_cell.border)
        target_cell.alignment = copy_style(source_cell.alignment)
        target_cell.protection = copy_style(source_cell.protection)
        target_cell.number_format = source_cell.number_format


def _sanitize_token(value: str, *, max_length: int) -> str:
    sanitized = str(value).replace("/", "_").replace("\\", "_")
    return sanitized[:max_length]


def _safe_group_name(keys) -> str:
    if not isinstance(keys, tuple):
        keys = (keys,)
    parts: list[str] = []
    for key in keys:
        text = str(key).strip()
        for char in '/\\:*?"<>|':
            text = text.replace(char, "_")
        parts.append(text)
    filename = "_".join(parts) if any(parts) else "(vide)"
    return filename[:200]


def _notify(callback: ProgressCallback | None, done: int, total: int, message: str) -> None:
    if callback is not None:
        callback(done, total, message)


def split_by_line_count(
    *,
    source_path: str,
    dataframe: pd.DataFrame,
    lines_per_file: int,
    output_dir: str | None,
    progress_callback: ProgressCallback | None = None,
) -> int:
    source = Path(source_path)
    destination = Path(output_dir) if output_dir else source.parent
    destination.mkdir(parents=True, exist_ok=True)
    is_csv = source.suffix.lower() == ".csv"

    if lines_per_file <= 0:
        raise ValueError("Le nombre de lignes par fichier doit être supérieur à 0.")

    total_rows = len(dataframe)
    if total_rows == 0:
        raise ValueError("Le fichier ne contient aucune donnée.")

    total_files = math.ceil(total_rows / lines_per_file)
    _notify(progress_callback, 0, total_files, "Préparation …")

    written = 0
    for i in range(total_files):
        chunk = dataframe.iloc[i * lines_per_file : (i + 1) * lines_per_file]
        safe_name = _sanitize_token(f"{source.stem}_part{i+1}", max_length=180)

        if is_csv:
            target = destination / f"{safe_name}.csv"
            chunk.to_csv(target, index=False, encoding="utf-8-sig")
        else:
            target = destination / f"{safe_name}.xlsx"
            chunk.to_excel(target, index=False, engine="xlsxwriter")

        written += 1
        _notify(progress_callback, written, total_files, f"{written} / {total_files} fichier(s) créé(s)")

    return written