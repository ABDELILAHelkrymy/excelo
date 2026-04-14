from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


EXCEL_SUFFIXES = {".xlsx", ".xls"}
CSV_SENTINEL = "(csv)"


@dataclass(frozen=True)
class SheetMetadata:
    row_count: int
    columns: tuple[str, ...]


def load_tabular_file(path: str | Path, sheet_name: str | int | None = 0) -> pd.DataFrame:
    file_path = Path(path)
    suffix = file_path.suffix.lower()

    if suffix in EXCEL_SUFFIXES:
        return pd.read_excel(file_path, sheet_name=sheet_name)

    if suffix == ".csv":
        for encoding in ("utf-8", "latin-1"):
            try:
                return pd.read_csv(file_path, encoding=encoding)
            except UnicodeDecodeError:
                continue
        return pd.read_csv(file_path, encoding="latin-1")

    raise ValueError(f"Unsupported file type: {file_path.suffix}")


def scan_tabular_source(path: str | Path) -> dict[str, SheetMetadata]:
    file_path = Path(path)
    suffix = file_path.suffix.lower()

    if suffix in EXCEL_SUFFIXES:
        return _scan_excel_source(file_path)
    if suffix == ".csv":
        return {CSV_SENTINEL: _scan_csv_source(file_path)}

    raise ValueError(f"Unsupported file type: {file_path.suffix}")


def _scan_excel_source(path: Path) -> dict[str, SheetMetadata]:
    if path.suffix.lower() == ".xls":
        return _scan_legacy_excel_source(path)

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        metadata: dict[str, SheetMetadata] = {}
        for worksheet in workbook.worksheets:
            header_cells = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            columns = tuple(str(value) for value in header_cells if value is not None)
            max_row = worksheet.max_row or 0
            row_count = max(max_row - 1, 0) if columns else max_row
            metadata[worksheet.title] = SheetMetadata(row_count=row_count, columns=columns)
        return metadata
    finally:
        workbook.close()


def _scan_legacy_excel_source(path: Path) -> dict[str, SheetMetadata]:
    metadata: dict[str, SheetMetadata] = {}
    with pd.ExcelFile(path) as excel_file:
        for sheet_name in excel_file.sheet_names:
            df = excel_file.parse(sheet_name)
            metadata[sheet_name] = SheetMetadata(
                row_count=len(df),
                columns=tuple(str(column) for column in df.columns),
            )
    return metadata


def _scan_csv_source(path: Path) -> SheetMetadata:
    columns: tuple[str, ...] | None = None
    for encoding in ("utf-8", "latin-1"):
        try:
            columns = tuple(str(column) for column in pd.read_csv(path, encoding=encoding, nrows=0).columns)
            break
        except UnicodeDecodeError:
            continue
    if columns is None:
        columns = tuple(str(column) for column in pd.read_csv(path, encoding="latin-1", nrows=0).columns)

    row_count = max(_count_newlines(path) - 1, 0)
    return SheetMetadata(row_count=row_count, columns=columns)


def _count_newlines(path: Path) -> int:
    total = 0
    with path.open("rb") as handle:
        while chunk := handle.read(1024 * 1024):
            total += chunk.count(b"\n")
    return total